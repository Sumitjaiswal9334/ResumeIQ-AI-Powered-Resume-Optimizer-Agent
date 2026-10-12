from __future__ import annotations

import json
import os
import re
import sqlite3
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch

DB_PATH = Path(__file__).with_name("resume_optimizer.db")

# Common job-related terms. The matching engine is deliberately transparent and editable.
KEYWORDS = [
    "python", "java", "javascript", "typescript", "c++", "sql", "nosql", "react", "node.js",
    "fastapi", "flask", "django", "rest api", "api", "git", "github", "docker", "kubernetes",
    "aws", "azure", "gcp", "linux", "ci/cd", "testing", "unit testing", "debugging",
    "data structures", "algorithms", "oop", "system design", "microservices", "postgresql",
    "mysql", "mongodb", "redis", "machine learning", "deep learning", "nlp", "llm", "rag",
    "langchain", "langgraph", "pytorch", "tensorflow", "scikit-learn", "pandas", "numpy",
    "data analysis", "statistics", "power bi", "tableau", "excel", "communication",
    "collaboration", "leadership", "problem solving", "agile", "scrum", "security",
    "performance optimization", "cloud", "data visualization", "generative ai", "prompt engineering",
    "vector database", "embeddings", "etl", "spark", "hadoop", "feature engineering",
]

def extract_resume_text(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(BytesIO(data))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
    elif ext == ".docx":
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Supported formats are PDF and DOCX.")
    if not text.strip():
        raise ValueError("No selectable text found. If this is a scanned PDF, run OCR first or paste the text manually.")
    return text.strip()

def _normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower().replace("–", "-").replace("—", "-"))

def _candidate_terms(job_description: str) -> list[str]:
    norm = _normalise(job_description)
    found = [k for k in KEYWORDS if re.search(r"(?<![a-z0-9])" + re.escape(k.lower()) + r"(?![a-z0-9])", norm)]
    # Add selected capitalized technical tokens / phrases from job description.
    phrases = re.findall(r"\b(?:[A-Z][A-Za-z0-9+#./-]{1,}|[A-Z][a-z]+\s(?:[A-Z][a-z]+|[A-Z]{2,}))\b", job_description)
    for phrase in phrases:
        p = phrase.strip().lower()
        if len(p) > 2 and p not in {"we", "you", "our", "the", "and", "with", "for"} and p not in found:
            if p in norm and len(found) < 45:
                found.append(p)
    return found[:45]

def match_resume_to_job(resume_text: str, job_description: str, job_title: str = "", company: str = "") -> dict[str, Any]:
    resume_norm = _normalise(resume_text)
    job_terms = _candidate_terms(job_description)
    title_terms = _candidate_terms(job_title)
    terms = list(dict.fromkeys(title_terms + job_terms))
    matched = [term for term in terms if re.search(r"(?<![a-z0-9])" + re.escape(term) + r"(?![a-z0-9])", resume_norm)]
    missing = [term for term in terms if term not in matched]
    score = round((len(matched) / len(terms)) * 100) if terms else 0
    return {
        "score": score,
        "job_keywords": terms,
        "matched_keywords": matched,
        "missing_keywords": missing,
        "method": "Exact normalized keyword coverage; synonyms and semantic equivalence are not inferred.",
    }

def _fallback_result(resume_text: str, match: dict[str, Any]) -> dict[str, Any]:
    return {
        "summary": "The resume has been compared with the job description using transparent keyword coverage. Review the potential gaps and tailor only the points that reflect your real experience.",
        "strengths": match["matched_keywords"][:8],
        "gaps": match["missing_keywords"][:10],
        "recommendations": [
            "Move the most relevant, genuinely possessed skills closer to the top of the resume.",
            "Use action verbs and explain the scope or outcome of projects where you have evidence.",
            "Include numbers only when you can substantiate them.",
            "Do not add a skill just because it appears in the job description.",
        ],
        "optimized_resume": resume_text,
        "cover_letter": "",
        "warnings": [],
    }

def analyze_resume(resume_text: str, job_description: str, job_title: str = "", company: str = "",
                   model: str = "gemini-3.5-flash-lite", candidate_name: str = "") -> dict[str, Any]:
    """Use Google Generative AI when configured; otherwise return a safe local fallback."""
    match = match_resume_to_job(resume_text, job_description, job_title, company)
    api_key = os.getenv("GOOGLE_API_KEY", "").strip()
    if not api_key:
        return _fallback_result(resume_text, match)
    from langchain_google_genai import ChatGoogleGenerativeAI
    from langchain_core.messages import SystemMessage, HumanMessage

    llm = ChatGoogleGenerativeAI(model=model, api_key=api_key, timeout=60, max_retries=1)
    prompt = f"""
Create a truthful resume review for the candidate and target role.
Return ONLY valid JSON with these keys:
summary (string), strengths (array of strings), gaps (array of strings),
recommendations (array of strings), optimized_resume (string), cover_letter (string),
warnings (array of strings).
Rules:
- Never invent employers, degrees, dates, certifications, skills, metrics, responsibilities, or achievements.
- Preserve factual details. If a detail is unclear, do not fabricate it; add a warning instead.
- You may reorder and rewrite existing content for clarity and relevance.
- Use plain text formatting for optimized_resume, with clear section headings.
- Cover letter must not claim experience not supported by the resume.
- Mention that recommendations need user verification if appropriate.
- Do not state that you calculated an official ATS score.
Candidate name: {candidate_name or "Not provided"}
Target job: {job_title or "Not provided"}
Company: {company or "Not provided"}
Resume text:
{resume_text[:18000]}
Job description:
{job_description[:12000]}
Keyword analysis:
{json.dumps(match, ensure_ascii=False)}
"""
    response = llm.invoke([
        SystemMessage(content="You are a careful resume editor. Be factual, conservative, and transparent."),
        HumanMessage(content=prompt),
    ])
    content = response.content
    if isinstance(content, list):
        content = "\n".join(str(x.get("text", x)) if isinstance(x, dict) else str(x) for x in content)
    content = str(content).strip()
    content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content, flags=re.I)
    try:
        result = json.loads(content)
        for key, default in {
            "summary": "", "strengths": [], "gaps": [], "recommendations": [],
            "optimized_resume": resume_text, "cover_letter": "", "warnings": []
        }.items():
            if key not in result or result[key] is None:
                result[key] = default
        return result
    except json.JSONDecodeError:
        return {
            **_fallback_result(resume_text, match),
            "warnings": ["The model returned a response that was not valid JSON. Showing the safe fallback instead."],
        }

def _connect():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE IF NOT EXISTS analyses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT NOT NULL,
            job_title TEXT,
            company TEXT,
            score INTEGER,
            summary TEXT,
            matched_count INTEGER,
            missing_count INTEGER
        )
    """)
    conn.commit()
    return conn

def save_analysis(result: dict[str, Any]) -> None:
    conn = _connect()
    try:
        conn.execute(
            "INSERT INTO analyses (created_at, job_title, company, score, summary, matched_count, missing_count) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                result.get("created_at", datetime.now().isoformat(timespec="seconds")),
                result.get("job_title", ""), result.get("company", ""),
                int(result.get("match", {}).get("score", 0)),
                result.get("ai", {}).get("summary", ""),
                len(result.get("match", {}).get("matched_keywords", [])),
                len(result.get("match", {}).get("missing_keywords", [])),
            ),
        )
        conn.commit()
    finally:
        conn.close()

def get_history(limit: int = 50) -> list[dict[str, Any]]:
    conn = _connect()
    try:
        rows = conn.execute("SELECT * FROM analyses ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()

def make_docx(text: str, candidate_name: str = "") -> bytes:
    doc = Document()
    doc.core_properties.title = "Optimized Resume"
    if candidate_name.strip():
        p = doc.add_heading(candidate_name.strip(), level=0)
    for line in text.splitlines():
        line = line.strip()
        if not line:
            doc.add_paragraph("")
        elif len(line) < 70 and (line.isupper() or line.endswith(":") or line.lower() in {
            "summary", "professional summary", "experience", "work experience", "education", "skills",
            "projects", "certifications", "achievements", "technical skills"
        }):
            doc.add_heading(line.rstrip(":"), level=1)
        else:
            doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

def make_pdf(text: str, candidate_name: str = "") -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.65*inch, leftMargin=0.65*inch,
                            topMargin=0.55*inch, bottomMargin=0.55*inch)
    styles = getSampleStyleSheet()
    story = []
    if candidate_name.strip():
        story.append(Paragraph(candidate_name.strip().replace("&", "&amp;"), styles["Title"]))
        story.append(Spacer(1, 8))
    for line in text.splitlines():
        safe = (line or " ").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if not line.strip():
            story.append(Spacer(1, 5))
        elif len(line.strip()) < 70 and (line.strip().isupper() or line.strip().endswith(":")):
            story.append(Paragraph(safe.rstrip(":"), styles["Heading2"]))
        else:
            story.append(Paragraph(safe, styles["BodyText"]))
            story.append(Spacer(1, 3))
    doc.build(story)
    return output.getvalue()
